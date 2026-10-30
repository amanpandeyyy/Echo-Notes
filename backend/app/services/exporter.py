import os
import io
import csv
import html
from datetime import datetime
from pathlib import Path
from typing import Optional, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
from app.config import settings
from app.models.schemas import Meeting, PersonaType

def clean_xml_text(text: Any) -> str:
    """Escapes HTML entities so ReportLab Paragraph never crashes on code or formulas."""
    if text is None:
        return ""
    return html.escape(str(text))

def generate_pdf_report(meeting: Meeting) -> str:
    """Generates a professional PDF report for a meeting using ReportLab."""
    pdf_filename = f"EchoNotes_{meeting.id}_{int(datetime.utcnow().timestamp())}.pdf"
    pdf_path = settings.EXPORTS_DIR / pdf_filename
    
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    primary_color = colors.HexColor("#0f172a") # Slate 900
    accent_color = colors.HexColor("#0284c7")  # Sky 600
    text_color = colors.HexColor("#334155")    # Slate 700
    bg_light = colors.HexColor("#f8fafc")      # Slate 50
    border_color = colors.HexColor("#e2e8f0")  # Slate 200
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=primary_color,
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=12
    )
    
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=accent_color,
        spaceBefore=12,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=text_color,
        spaceAfter=4
    )
    
    bullet_style = ParagraphStyle(
        'Bullet_Custom',
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-12,
        spaceAfter=3
    )
    
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=text_color
    )
    
    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.white
    )

    story = []
    
    # Header Branding
    story.append(Paragraph("ECHO-NOTES • MEETING INTELLIGENCE REPORT", subtitle_style))
    story.append(Paragraph(clean_xml_text(meeting.title), title_style))
    
    # Metadata bar
    meta_text = f"<b>Date:</b> {meeting.created_at.strftime('%Y-%m-%d %H:%M UTC')} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Duration:</b> {meeting.duration_str} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Language:</b> {clean_xml_text(meeting.language.upper())} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Persona:</b> {clean_xml_text(meeting.persona.value.replace('_', ' ').title())}"
    story.append(Paragraph(meta_text, subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=border_color, spaceAfter=10))
    
    # Summary Section
    if meeting.summary:
        story.append(Paragraph("Executive Summary", h1_style))
        story.append(Paragraph(clean_xml_text(meeting.summary.executive_summary), body_style))
        story.append(Spacer(1, 6))
        
        if meeting.summary.bullet_points:
            story.append(Paragraph("Key Takeaways", h1_style))
            for bp in meeting.summary.bullet_points:
                story.append(Paragraph(f"• {clean_xml_text(bp)}", bullet_style))
            story.append(Spacer(1, 6))
            
        # Decisions Table
        if meeting.summary.decisions:
            story.append(Paragraph("Decisions Log", h1_style))
            dec_data = [[Paragraph("Decision", table_header_style), Paragraph("Rationale", table_header_style), Paragraph("Made By", table_header_style)]]
            for d in meeting.summary.decisions:
                dec_data.append([
                    Paragraph(clean_xml_text(d.decision), table_cell_style),
                    Paragraph(clean_xml_text(d.rationale or "—"), table_cell_style),
                    Paragraph(clean_xml_text(d.made_by or "Team"), table_cell_style)
                ])
            
            t_dec = Table(dec_data, colWidths=[200, 220, 120])
            t_dec.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), primary_color),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('GRID', (0,0), (-1,-1), 0.5, border_color),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ]))
            story.append(t_dec)
            story.append(Spacer(1, 8))
            
        # Action Items Table
        if meeting.summary.action_items:
            story.append(Paragraph("Action Items & Deliverables", h1_style))
            act_data = [[
                Paragraph("Status", table_header_style),
                Paragraph("Task", table_header_style),
                Paragraph("Owner", table_header_style),
                Paragraph("Priority", table_header_style),
                Paragraph("Due Date", table_header_style)
            ]]
            for a in meeting.summary.action_items:
                status_str = "✓ Done" if a.completed else "Pending"
                act_data.append([
                    Paragraph(status_str, table_cell_style),
                    Paragraph(clean_xml_text(a.task), table_cell_style),
                    Paragraph(clean_xml_text(a.owner or "Unassigned"), table_cell_style),
                    Paragraph(clean_xml_text(a.priority.value.upper()), table_cell_style),
                    Paragraph(clean_xml_text(a.due_date or "—"), table_cell_style)
                ])
                
            t_act = Table(act_data, colWidths=[60, 220, 100, 70, 90])
            t_act.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), accent_color),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('GRID', (0,0), (-1,-1), 0.5, border_color),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ]))
            story.append(t_act)
            story.append(Spacer(1, 8))
            
        # Open Questions
        if meeting.summary.open_questions:
            story.append(Paragraph("Open Questions & Follow-ups", h1_style))
            for q in meeting.summary.open_questions:
                story.append(Paragraph(f"❓ <b>{clean_xml_text(q.question)}</b>", bullet_style))
                if q.context:
                    story.append(Paragraph(f"<i>Context: {clean_xml_text(q.context)}</i>", ParagraphStyle('SubContext', parent=body_style, leftIndent=16, fontSize=8.5, textColor=colors.HexColor("#64748b"))))
            story.append(Spacer(1, 8))

        # Persona-Specific Intelligence Section
        if meeting.summary.persona_insights:
            pi = meeting.summary.persona_insights
            
            # Academic Lecture & Exam Study Notes
            if pi.academic_notes:
                an = pi.academic_notes
                story.append(Paragraph("Academic Lecture Intelligence & Study Notes", h1_style))
                if an.get("subject"):
                    story.append(Paragraph(f"<b>Course / Subject:</b> {clean_xml_text(an['subject'])}", body_style))
                if an.get("handwritten_summary"):
                    story.append(Paragraph("<b>Student Lecture Summary:</b>", body_style))
                    story.append(Paragraph(clean_xml_text(an["handwritten_summary"]), body_style))
                    story.append(Spacer(1, 4))
                
                # High-Yield Important Points
                if an.get("important_points"):
                    story.append(Paragraph("High-Yield Exam Points & Realizations", h1_style))
                    for pt in an["important_points"]:
                        story.append(Paragraph(f"• {clean_xml_text(pt)}", bullet_style))
                    story.append(Spacer(1, 4))
                    
                # Key Definitions
                if an.get("key_definitions"):
                    story.append(Paragraph("Key Definitions & Theoretical Terms", h1_style))
                    for kd in an["key_definitions"]:
                        term = clean_xml_text(kd.get("term", ""))
                        defn = clean_xml_text(kd.get("definition", ""))
                        story.append(Paragraph(f"<b>{term}:</b> {defn}", bullet_style))
                    story.append(Spacer(1, 4))
                    
                # Practice Questions
                if an.get("practice_questions"):
                    story.append(Paragraph("Practice Questions & Explanations", h1_style))
                    for pq in an["practice_questions"]:
                        q_txt = clean_xml_text(pq.get("q", ""))
                        a_txt = clean_xml_text(pq.get("a", ""))
                        story.append(Paragraph(f"<b>Q: {q_txt}</b>", bullet_style))
                        if a_txt:
                            story.append(Paragraph(f"<i>Answer: {a_txt}</i>", ParagraphStyle('AnsStyle', parent=body_style, leftIndent=16, fontSize=8.5, textColor=colors.HexColor("#334155"))))
                    story.append(Spacer(1, 4))
                    
                # Study Checklist
                if an.get("study_checklist"):
                    story.append(Paragraph("Study Checklist & Review Items", h1_style))
                    for sc in an["study_checklist"]:
                        story.append(Paragraph(f"☐ {clean_xml_text(sc.get('text', ''))}", bullet_style))
                    story.append(Spacer(1, 6))

            # Recruiter Scorecard
            if pi.scorecard:
                sc = pi.scorecard
                story.append(Paragraph("Recruiter Candidate Evaluation Scorecard", h1_style))
                sc_data = [[Paragraph("Dimension", table_header_style), Paragraph("Evaluation & Notes", table_header_style)]]
                for key, val in sc.items():
                    label = key.replace("_", " ").title()
                    sc_data.append([
                        Paragraph(f"<b>{clean_xml_text(label)}</b>", table_cell_style),
                        Paragraph(clean_xml_text(val), table_cell_style)
                    ])
                t_sc = Table(sc_data, colWidths=[160, 380])
                t_sc.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), primary_color),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('GRID', (0,0), (-1,-1), 0.5, border_color),
                    ('TOPPADDING', (0,0), (-1,-1), 4),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ]))
                story.append(t_sc)
                story.append(Spacer(1, 8))

            # PM Sprint & Delivery Tracking
            if pi.pm_tracking:
                pm = pi.pm_tracking
                story.append(Paragraph("Project Manager Sprint & Delivery Tracking", h1_style))
                pm_data = [[Paragraph("Metric / Area", table_header_style), Paragraph("Status / Assessment", table_header_style)]]
                for key, val in pm.items():
                    label = key.replace("_", " ").title()
                    val_str = f"{val}%" if key == "milestone_completion_pct" else str(val)
                    pm_data.append([
                        Paragraph(f"<b>{clean_xml_text(label)}</b>", table_cell_style),
                        Paragraph(clean_xml_text(val_str), table_cell_style)
                    ])
                t_pm = Table(pm_data, colWidths=[160, 380])
                t_pm.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), accent_color),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('GRID', (0,0), (-1,-1), 0.5, border_color),
                    ('TOPPADDING', (0,0), (-1,-1), 4),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ]))
                story.append(t_pm)
                story.append(Spacer(1, 8))

    # Transcript Section
    if meeting.transcript:
        story.append(Paragraph("Full Timestamped Transcript", h1_style))
        for seg in meeting.transcript:
            seg_text = f"<b>[{clean_xml_text(seg.start_str)}] {clean_xml_text(seg.speaker)}:</b> {clean_xml_text(seg.text)}"
            story.append(Paragraph(seg_text, body_style))
            
    doc.build(story)
    return str(pdf_path)

def generate_docx_report(meeting: Meeting) -> str:
    """Generates a styled DOCX report for a meeting using python-docx."""
    docx_filename = f"EchoNotes_{meeting.id}_{int(datetime.utcnow().timestamp())}.docx"
    docx_path = settings.EXPORTS_DIR / docx_filename
    
    doc = Document()
    
    # Title & Metadata
    title_p = doc.add_heading(meeting.title, level=0)
    title_p.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Generated by Echo-Notes AI • Date: {meeting.created_at.strftime('%Y-%m-%d %H:%M UTC')} • Duration: {meeting.duration_str} • Language: {meeting.language.upper()}\n").italic = True
    
    if meeting.summary:
        # Executive Summary
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph(meeting.summary.executive_summary)
        
        # Key Takeaways
        if meeting.summary.bullet_points:
            doc.add_heading('Key Takeaways', level=1)
            for bp in meeting.summary.bullet_points:
                doc.add_paragraph(bp, style='List Bullet')
                
        # Decisions Log
        if meeting.summary.decisions:
            doc.add_heading('Decisions Log', level=1)
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Light Shading Accent 1'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Decision'
            hdr_cells[1].text = 'Rationale'
            hdr_cells[2].text = 'Made By'
            for d in meeting.summary.decisions:
                row_cells = table.add_row().cells
                row_cells[0].text = d.decision
                row_cells[1].text = d.rationale or ''
                row_cells[2].text = d.made_by or ''
            doc.add_paragraph()
            
        # Action Items
        if meeting.summary.action_items:
            doc.add_heading('Action Items', level=1)
            table = doc.add_table(rows=1, cols=5)
            table.style = 'Light Shading Accent 1'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Status'
            hdr_cells[1].text = 'Task'
            hdr_cells[2].text = 'Owner'
            hdr_cells[3].text = 'Priority'
            hdr_cells[4].text = 'Due Date'
            for a in meeting.summary.action_items:
                row_cells = table.add_row().cells
                row_cells[0].text = 'Done' if a.completed else 'Pending'
                row_cells[1].text = a.task
                row_cells[2].text = a.owner or 'Unassigned'
                row_cells[3].text = a.priority.value.upper()
                row_cells[4].text = a.due_date or ''
            doc.add_paragraph()
            
        # Open Questions
        if meeting.summary.open_questions:
            doc.add_heading('Open Questions', level=1)
            for q in meeting.summary.open_questions:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(q.question).bold = True
                if q.context:
                    p.add_run(f" (Context: {q.context})").italic = True

        # Persona-Specific Intelligence Section
        if meeting.summary.persona_insights:
            pi = meeting.summary.persona_insights

            # Academic Lecture & Exam Study Notes
            if pi.academic_notes:
                an = pi.academic_notes
                doc.add_heading('Academic Study Notes & Lecture Intelligence', level=1)
                if an.get("subject"):
                    p = doc.add_paragraph()
                    p.add_run("Course / Subject: ").bold = True
                    p.add_run(str(an["subject"]))
                if an.get("handwritten_summary"):
                    doc.add_heading('Student Lecture Summary', level=2)
                    doc.add_paragraph(str(an["handwritten_summary"]))
                if an.get("important_points"):
                    doc.add_heading('High-Yield Exam Points & Key Realizations', level=2)
                    for pt in an["important_points"]:
                        doc.add_paragraph(str(pt), style='List Bullet')
                if an.get("key_definitions"):
                    doc.add_heading('Key Definitions & Theoretical Terms', level=2)
                    for kd in an["key_definitions"]:
                        p = doc.add_paragraph(style='List Bullet')
                        p.add_run(f"{kd.get('term', '')}: ").bold = True
                        p.add_run(str(kd.get('definition', '')))
                if an.get("practice_questions"):
                    doc.add_heading('Practice Questions & Explanations', level=2)
                    for pq in an["practice_questions"]:
                        p = doc.add_paragraph()
                        p.add_run(f"Q: {pq.get('q', '')}").bold = True
                        if pq.get('a'):
                            p2 = doc.add_paragraph()
                            p2.add_run(f"Answer: {pq.get('a', '')}").italic = True
                if an.get("study_checklist"):
                    doc.add_heading('Study Checklist', level=2)
                    for sc in an["study_checklist"]:
                        doc.add_paragraph(f"☐ {sc.get('text', '')}", style='List Bullet')

            # Recruiter Scorecard
            if pi.scorecard:
                sc = pi.scorecard
                doc.add_heading('Recruiter Candidate Evaluation Scorecard', level=1)
                table = doc.add_table(rows=1, cols=2)
                table.style = 'Light Shading Accent 1'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Dimension'
                hdr_cells[1].text = 'Evaluation'
                for k, v in sc.items():
                    row_cells = table.add_row().cells
                    row_cells[0].text = k.replace('_', ' ').title()
                    row_cells[1].text = str(v)
                doc.add_paragraph()

            # PM Tracking
            if pi.pm_tracking:
                pm = pi.pm_tracking
                doc.add_heading('Project Manager Sprint & Delivery Tracking', level=1)
                table = doc.add_table(rows=1, cols=2)
                table.style = 'Light Shading Accent 1'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Metric / Area'
                hdr_cells[1].text = 'Status'
                for k, v in pm.items():
                    row_cells = table.add_row().cells
                    row_cells[0].text = k.replace('_', ' ').title()
                    row_cells[1].text = f"{v}%" if k == "milestone_completion_pct" else str(v)
                doc.add_paragraph()

    # Transcript
    if meeting.transcript:
        doc.add_heading('Transcript', level=1)
        for seg in meeting.transcript:
            p = doc.add_paragraph()
            p.add_run(f"[{seg.start_str}] {seg.speaker}: ").bold = True
            p.add_run(seg.text)
            
    doc.save(str(docx_path))
    return str(docx_path)

def generate_csv_report(meeting: Meeting) -> str:
    """Generates a CSV export of action items, decisions, persona insights, and transcript."""
    csv_filename = f"EchoNotes_{meeting.id}_{int(datetime.utcnow().timestamp())}.csv"
    csv_path = settings.EXPORTS_DIR / csv_filename
    
    with open(csv_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        
        # Section 1: Metadata
        writer.writerow(["MEETING METADATA"])
        writer.writerow(["Title", meeting.title])
        writer.writerow(["Date", meeting.created_at.isoformat()])
        writer.writerow(["Duration", meeting.duration_str])
        writer.writerow(["Language", meeting.language])
        writer.writerow(["Persona", meeting.persona.value])
        writer.writerow([])
        
        # Section 2: Action Items
        writer.writerow(["ACTION ITEMS"])
        writer.writerow(["Status", "Task", "Owner", "Priority", "Due Date"])
        if meeting.summary and meeting.summary.action_items:
            for a in meeting.summary.action_items:
                writer.writerow(["Done" if a.completed else "Pending", a.task, a.owner, a.priority.value, a.due_date])
        writer.writerow([])
        
        # Section 3: Decisions
        writer.writerow(["DECISIONS LOG"])
        writer.writerow(["Decision", "Rationale", "Made By"])
        if meeting.summary and meeting.summary.decisions:
            for d in meeting.summary.decisions:
                writer.writerow([d.decision, d.rationale, d.made_by])
        writer.writerow([])

        # Section 4: Persona Insights (Academic / Recruiter / PM)
        if meeting.summary and meeting.summary.persona_insights:
            pi = meeting.summary.persona_insights
            if pi.academic_notes:
                an = pi.academic_notes
                writer.writerow(["ACADEMIC LECTURE NOTES"])
                if an.get("subject"):
                    writer.writerow(["Subject", an["subject"]])
                if an.get("important_points"):
                    writer.writerow(["Important Points"])
                    for pt in an["important_points"]:
                        writer.writerow(["Point", pt])
                if an.get("key_definitions"):
                    writer.writerow(["Key Definitions"])
                    for kd in an["key_definitions"]:
                        writer.writerow([kd.get("term"), kd.get("definition")])
                writer.writerow([])
            if pi.scorecard:
                writer.writerow(["RECRUITER SCORECARD"])
                for k, v in pi.scorecard.items():
                    writer.writerow([k.replace("_", " ").title(), str(v)])
                writer.writerow([])
            if pi.pm_tracking:
                writer.writerow(["PROJECT MANAGER TRACKING"])
                for k, v in pi.pm_tracking.items():
                    writer.writerow([k.replace("_", " ").title(), str(v)])
                writer.writerow([])
        
        # Section 5: Transcript
        writer.writerow(["TRANSCRIPT SEGMENTS"])
        writer.writerow(["Segment ID", "Timestamp", "Speaker", "Text"])
        for seg in meeting.transcript:
            writer.writerow([seg.id, seg.start_str, seg.speaker, seg.text])
            
    return str(csv_path)

